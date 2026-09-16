"""
One HTML document, and the PDF, Word file and JSON derived from it.

The lab's whole report design is the one Django template, rendered server side, so the preview the
scientist reads, the PDF filed with the case and the Word file they edit cannot disagree. The JSON
comes off the same context, so the structured record and the human document describe the same
classifications at the same versions - built by whichever app owns that template's shape, else as
the canonical context dump.

Entry point: render_case_report(template, context, title) -> RenderedCaseReport.
"""
import io
import json
import re
from dataclasses import dataclass
from typing import Optional

from django.core.serializers.json import DjangoJSONEncoder
from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor
from html2docx import html2docx

from classification.models.classification_report_models import (
    ClassificationReportTemplate,
    case_report_json_signal,
)
from classification.report.template_validation import html_to_pdf, render_html


@dataclass
class RenderedCaseReport:
    html: str
    pdf: bytes
    docx: bytes
    json_output: dict
    # The context as stored - JSON safe, so the same documents can be re-rendered after a template fix
    context_snapshot: dict


def json_safe(context: dict) -> dict:
    """ Dates and decimals the way a JSONField will take them back """
    return json.loads(json.dumps(context, cls=DjangoJSONEncoder))


def render_json(report_template: ClassificationReportTemplate, context: dict) -> dict:
    """ The report's structured record: the app that owns this template's shape, else the canonical
        context dump - so there is always one.

        A JSON another system parses is an interface, and keeping it in step with that system is
        code rather than config a lab edits, so it is written in Python by the app that has to.
        Anything raised there fails the build rather than quietly falling through to a differently
        shaped document. """
    for _receiver, result in case_report_json_signal.send(sender=ClassificationReportTemplate,
                                                          report_template=report_template,
                                                          context=context):
        if result is not None:
            return result
    return json_safe(context)


# html2docx has no notion of a document head: it prints the contents of <style> and <script> as the
# first paragraph of the Word file. They are the print design, not the report, so they come out first.
# The DRAFT watermark goes with them - it is a positioned, rotated element with no Word equivalent,
# and the Word file gets a banner paragraph instead
NON_CONTENT_PATTERN = re.compile(r"<(head|style|script)\b.*?</\1\s*>", re.DOTALL | re.IGNORECASE)
WATERMARK_PATTERN = re.compile(r'<(\w+)[^>]*\bclass="[^"]*\bdraft-watermark\b[^"]*"[^>]*>.*?</\1\s*>',
                               re.DOTALL | re.IGNORECASE)

# A <pre> block reaches Word with its line breaks and its spaces but in the body's proportional
# font, so a column aligned Results Summary lines up nowhere. <code> is the one tag html2docx takes
# a font from, and it has to stay open to the </pre>: closing it drops the run html2docx turns the
# block's newlines into Word breaks on, and the whole block arrives as one line
PRE_OPEN_PATTERN = re.compile(r"<pre\b[^>]*>", re.IGNORECASE)
HTML2DOCX_MONO_FONT = "Mono"  # what html2docx writes for <code>, and not a font Word has
PRE_FONT_NAME = "Courier New"
PRE_FONT_SIZE = Pt(10)

DRAFT_BANNER_TEXT = "DRAFT - not the issued report"
DRAFT_BANNER_COLOUR = RGBColor(0xC0, 0x00, 0x00)
DRAFT_BANNER_SIZE = Pt(14)


def _monospace_pre_blocks(document: DocxDocument):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name == HTML2DOCX_MONO_FONT:
                run.font.name = PRE_FONT_NAME
                run.font.size = PRE_FONT_SIZE


def _add_draft_banner(document: DocxDocument):
    """ The Word file's answer to the PDF's watermark - first thing on the page, so a printed draft
        cannot be mistaken for the copy that went out with the case """
    first = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
    run = first.insert_paragraph_before().add_run(DRAFT_BANNER_TEXT)
    run.bold = True
    run.font.color.rgb = DRAFT_BANNER_COLOUR
    run.font.size = DRAFT_BANNER_SIZE


def render_docx(html: str, title: str, draft: bool = False) -> bytes:
    content = WATERMARK_PATTERN.sub("", NON_CONTENT_PATTERN.sub("", html))
    document = Document(html2docx(PRE_OPEN_PATTERN.sub(r"\g<0><code>", content), title=title))
    _monospace_pre_blocks(document)
    if draft:
        _add_draft_banner(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_case_report(template: ClassificationReportTemplate, context: dict,
                       title: Optional[str] = None) -> RenderedCaseReport:
    html = render_html(template.case_template, context)
    return RenderedCaseReport(
        html=html,
        pdf=html_to_pdf(html),
        docx=render_docx(html, title or template.name, draft=bool(context.get("draft"))),
        json_output=render_json(template, context),
        context_snapshot=json_safe(context),
    )
